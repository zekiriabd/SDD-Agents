#!/usr/bin/env python3
"""Génère les outils d'accès aux sources déclarées — 0 token, déterministe.

Ce que ce script défend : *la chaîne « source déclarée → schéma figé → outil →
contrat » est régénérable, donc vérifiable.* Un outil d'accès écrit à la main
diverge de sa déclaration en trois semaines, et personne ne le voit : le code
filtre sur un champ que la déclaration n'expose pas, la description que lit le
modèle n'est plus celle qu'un humain a revue, et le schéma n'est plus celui
contre lequel le démarrage valide. Ici, une déclaration produit toujours le
même code, et `--check` dit quand les deux ont divergé.

Trois modes, trois moments :

    --infer --source ID   une fois, au début : observe la donnée réelle et
                          écrit un brouillon de schéma que l'on RELIT.
                          N'écrase jamais un schéma existant sans --force.
    --write               à chaque changement de déclaration : réécrit les
                          wrappers Python, crée les contrats manquants.
                          Ne touche jamais à un contrat existant (il est
                          complété par architect-tools et dev-prompt).
    --check               en CI, le défaut : régénère en mémoire et compare.
                          Exit 1 si le code a dérivé de la déclaration.

Il n'appelle aucun réseau et n'exécute pas l'application : une source distante
(`http-api`, `mcp`) ou un store distant ne s'infère donc pas ici — il faut
`--from-sample`, un fichier de réponse capturé. C'est volontaire : un
générateur qui joint un serveur ne tourne pas en CI, et une gate qu'on ne peut
pas jouer est une gate qu'on saute.

Usage :
    python .sdda/sdda.py gen-source-tools --check --json
    python .sdda/sdda.py gen-source-tools --infer --source order_tracking
    python .sdda/sdda.py gen-source-tools --infer --source crm_customer --from-sample sample.json
    python .sdda/sdda.py gen-source-tools --write --mission 1
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import re
import sys
from pathlib import Path
from typing import Any, Iterator

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from sdda_lib import markdown_io, paths, schema_infer, source_registry as sr  # noqa: E402
from sdda_lib.errors import Report  # noqa: E402
from sdda_lib.layered_config import active_stacks, read_project_section, read_stack_section_kv  # noqa: E402
from sdda_scripts._common import add_common_args, ensure_utf8_stdout, finish, resolve_root  # noqa: E402

DECLARED = "declared-sources"
BANNER = "GÉNÉRÉ par gen_source_tools.py — NE PAS ÉDITER."

#: JSON Schema -> annotation Python du modèle généré.
PY_TYPES = {
    "string": "str", "integer": "int", "number": "float", "boolean": "bool",
    "object": "dict[str, Any]", "array": "list[Any]", "null": "None",
}

#: Erreurs déclarées par type d'outil : code -> (signification, comportement attendu).
DECLARED_ERRORS: dict[str, dict[str, tuple[str, str]]] = {
    "lookup": {
        "NOT_FOUND": ("aucun enregistrement pour cette clé", "le dire à l'utilisateur, ne pas réessayer"),
        "SOURCE_STALE": ("l'instantané dépasse max_staleness_hours", "répondre en signalant explicitement la date de la donnée"),
        "SOURCE_UNAVAILABLE": ("source illisible (fichier réécrit, hôte injoignable)", "un seul retry, puis échec explicite"),
        "TIMEOUT": ("budget de lecture dépassé", "échec explicite, jamais un résultat partiel muet"),
    },
    "search": {
        "TOO_MANY_RECORDS": ("plus d'enregistrements que le plafond", "affiner les filtres avant de répondre, ne jamais conclure sur le tronqué"),
        "INVALID_FILTER": ("filtre sur un champ non exposé, ou valeur hors domaine", "corriger le filtre, ne pas contourner par un autre outil"),
        "SOURCE_STALE": ("l'instantané dépasse max_staleness_hours", "répondre en signalant explicitement la date de la donnée"),
        "SOURCE_UNAVAILABLE": ("source illisible (fichier réécrit, hôte injoignable)", "un seul retry, puis échec explicite"),
        "TIMEOUT": ("budget de lecture dépassé", "échec explicite, jamais un résultat partiel muet"),
    },
    "count": {
        "INVALID_FILTER": ("filtre sur un champ non exposé, ou valeur hors domaine", "corriger le filtre, ne pas contourner par un autre outil"),
        "SOURCE_STALE": ("l'instantané dépasse max_staleness_hours", "répondre en signalant explicitement la date de la donnée"),
        "SOURCE_UNAVAILABLE": ("source illisible (fichier réécrit, hôte injoignable)", "un seul retry, puis échec explicite"),
        "TIMEOUT": ("budget de lecture dépassé", "échec explicite, jamais un résultat partiel muet"),
    },
}


# ---------------------------------------------------------------------------
# Contexte
# ---------------------------------------------------------------------------
class Context:
    """Ce qu'il faut savoir du projet pour générer, résolu une seule fois."""

    def __init__(self, root: Path, *, mission: str | None = None, src_root: Path | None = None):
        self.root = root
        self.section = read_stack_section_kv(root, "Active Data Sources")
        self.registry = sr.load_registry(root, self.section)
        self.app = str(read_project_section(root).get("AppName") or "App").strip() or "App"
        self.mission = mission or self._detect_mission()
        self.mission_name = self._mission_name()
        self.src_root = src_root or paths.app_src_root(root, self.app)

    def _detect_mission(self) -> str:
        found = sorted(p.name.split("-", 1)[0] for p in paths.missions_dir(self.root).glob("*-*.md"))
        uniques = sorted(set(found))
        return uniques[0] if len(uniques) == 1 else ""

    def _mission_name(self) -> str:
        if not self.mission:
            return ""
        for p in sorted(paths.missions_dir(self.root).glob(f"{self.mission}-*.md")):
            return p.name[: -len(".md")]
        return ""

    def tools_dir(self) -> Path:
        return self.src_root / "data" / "tools"

    def contract_path(self, source_id: str, kind: str) -> Path:
        return paths.contracts_dir(self.root, "tools") / f"{self.contract_id(source_id, kind)}.tool.md"

    def contract_id(self, source_id: str, kind: str) -> str:
        return f"{self.mission}-{source_id.replace('_', '-')}-{kind}"

    def wrapper_path(self, source_id: str, kind: str) -> Path:
        return self.tools_dir() / f"{source_id}_{kind}.py"

    def schema_path(self, source_id: str) -> Path:
        # À côté des wrappers, dans le paquet : c'est là que `schema_guard`
        # le cherche au démarrage. Suit `src_root` quand il est redirigé.
        return self.src_root / "data" / "schemas" / f"{source_id}.schema.json"

    def envelope_int(self, key: str, default: int) -> int:
        try:
            value = int(self.section.get(key))
        except (TypeError, ValueError):
            return default
        return value if value > 0 else default


# ---------------------------------------------------------------------------
# Lecture des enregistrements — fichiers locaux uniquement, aucun réseau
# ---------------------------------------------------------------------------
def _walk(payload: Any, dotted: str) -> Any:
    for part in dotted.split("."):
        if not isinstance(payload, dict) or part not in payload:
            raise ValueError(f"`records_path: {dotted}` : segment `{part}` absent de la réponse")
        payload = payload[part]
    return payload


def read_records(ctx: Context, source_id: str, src: dict[str, Any], report: Report,
                 sample: Path | None = None, limit: int = 200_000) -> tuple[list[dict[str, Any]], bool]:
    """Les enregistrements observables d'une source. `(records, coerce)`.

    `coerce` dit à l'inférence que les valeurs sont des cellules texte et
    doivent être typées — c'est vrai des CSV et des XLSX, et faux partout
    ailleurs. Le distinguer ici évite de retyper un JSON qui portait déjà ses
    types, ce qui transformerait `"0012345"` en `12345`.
    """
    if sample is not None:
        return _read_sample(sample, src, report), False

    connector = str(src.get("connector") or "")
    if connector != "file":
        report.error(
            "DATA_SCHEMA_SAMPLE_REQUIRED",
            f"source `{source_id}` ({connector}) : l'inférence exigerait un appel réseau",
            fix="capturer une réponse représentative dans un fichier JSON et relancer avec "
                "`--from-sample <fichier>` — ce script ne joint jamais un hôte, sinon il ne tournerait pas en CI",
        )
        return [], False

    store = ctx.registry.stores.get(str(src.get("store") or ""))
    if store is None:
        report.error("DATA_STORE_UNKNOWN", f"source `{source_id}` : store `{src.get('store')}` non déclaré",
                     fix="déclarer le store, ou corriger la référence")
        return [], False
    kind = str(store.get("kind") or "")
    if kind not in sr.LOCAL_KINDS:
        report.error(
            "DATA_SCHEMA_SAMPLE_REQUIRED",
            f"source `{source_id}` : le store `{store.get('id')}` est distant (`{kind}`)",
            fix="synchroniser un échantillon en local, ou relancer avec `--from-sample <fichier>`",
        )
        return [], False

    files = _resolve_files(ctx.root, store, src, source_id, report)
    fmt = str(src.get("format") or "").strip().lower()
    encoding = str(src.get("encoding") or "utf-8")
    records: list[dict[str, Any]] = []
    for path in files:
        if len(records) >= limit:
            break
        try:
            records.extend(_read_file(path, fmt, src, encoding))
        except (OSError, ValueError, UnicodeDecodeError) as exc:
            report.error(
                "DATA_SOURCE_UNREADABLE",
                f"source `{source_id}` : `{paths.rel(ctx.root, path)}` illisible — {exc}",
                fix="vérifier le format, le délimiteur et l'encodage déclarés ; `errors=strict` est "
                    "délibéré : un remplacement silencieux produirait des données fausses",
            )
    return records[:limit], fmt in ("csv", "tsv", "xlsx")


def _resolve_files(root: Path, store: dict[str, Any], src: dict[str, Any], source_id: str,
                   report: Report) -> list[Path]:
    declared = str(store.get("root") or "")
    candidate = Path(declared)
    base = candidate if (candidate.is_absolute() or declared.startswith(("//", "\\\\"))) else (root / candidate)
    try:
        store_root = base.resolve()
    except OSError:
        store_root = base
    if not store_root.is_dir():
        report.error("DATA_STORE_UNREACHABLE", f"source `{source_id}` : racine `{declared}` introuvable",
                     fix="monter le partage ou corriger `root`")
        return []
    glob = str(src.get("glob") or "")
    try:
        matches = sorted(p for p in store_root.glob(glob)
                         if p.is_file() and not p.is_symlink() and p.resolve().is_relative_to(store_root))
    except (OSError, ValueError):
        matches = []
    if not matches:
        report.error("DATA_SOURCE_EMPTY", f"source `{source_id}` : `{glob}` ne résout aucun fichier",
                     fix="déposer les exports, ou corriger le glob — on n'infère pas un schéma depuis rien")
    return matches


def _read_file(path: Path, fmt: str, src: dict[str, Any], encoding: str) -> Iterator[dict[str, Any]]:
    if fmt == "jsonl":
        with path.open("r", encoding=encoding, errors="strict") as fh:
            for line in fh:
                if line.strip():
                    value = json.loads(line)
                    if isinstance(value, dict):
                        yield value
        return
    if fmt in ("array", "object"):
        payload = json.loads(path.read_text(encoding=encoding, errors="strict"))
        records_path = str(src.get("records_path") or "").strip()
        if records_path:
            payload = _walk(payload, records_path)
        if isinstance(payload, list):
            yield from (r for r in payload if isinstance(r, dict))
        elif isinstance(payload, dict):
            yield payload
        return
    if fmt in ("csv", "tsv"):
        delimiter = str(src.get("delimiter") or ("\t" if fmt == "tsv" else ","))
        header_row = int(src.get("header_row") or 1)
        text = path.read_text(encoding=encoding, errors="strict")
        lines = text.split("\n")[header_row - 1:]
        reader = csv.DictReader(io.StringIO("\n".join(lines)), delimiter=delimiter)
        for row in reader:
            yield {k: v for k, v in row.items() if k is not None}
        return
    if fmt == "xlsx":
        yield from _read_xlsx(path, src)
        return
    if fmt == "parquet":
        yield from _read_parquet(path)
        return
    raise ValueError(f"format `{fmt}` sans lecteur")


def _read_xlsx(path: Path, src: dict[str, Any]) -> Iterator[dict[str, Any]]:
    try:
        from openpyxl import load_workbook            # type: ignore[import-not-found]
    except ImportError as exc:                        # pragma: no cover — dépend de l'environnement
        raise ValueError("`openpyxl` absent : l'épingler dans le .libs.json du framework actif") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_name = str(src.get("sheet") or "")
    sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
    header_row = int(src.get("header_row") or 1)
    header: list[str] = []
    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if index < header_row:
            continue
        if index == header_row:
            header = [str(c) if c is not None else f"col_{i}" for i, c in enumerate(row)]
            continue
        yield {header[i]: ("" if c is None else str(c)) for i, c in enumerate(row) if i < len(header)}


def _read_parquet(path: Path) -> Iterator[dict[str, Any]]:
    try:
        import pyarrow.parquet as pq                  # type: ignore[import-not-found]
    except ImportError as exc:                        # pragma: no cover — dépend de l'environnement
        raise ValueError("`pyarrow` absent : l'épingler dans le .libs.json du framework actif") from exc
    yield from pq.read_table(path).to_pylist()


def _read_sample(sample: Path, src: dict[str, Any], report: Report) -> list[dict[str, Any]]:
    try:
        payload = json.loads(markdown_io.read_text(sample))
    except (OSError, ValueError) as exc:
        report.error("DATA_SCHEMA_SAMPLE_INVALID", f"échantillon `{sample}` illisible — {exc}",
                     fix="fournir un JSON : soit la réponse brute, soit un tableau d'enregistrements")
        return []
    records_path = str(src.get("records_path") or "").strip()
    if records_path and isinstance(payload, dict):
        try:
            payload = _walk(payload, records_path)
        except ValueError as exc:
            report.error("DATA_SCHEMA_SAMPLE_INVALID", f"échantillon `{sample}` : {exc}",
                         fix="l'échantillon doit avoir la même forme que la réponse réelle")
            return []
    if isinstance(payload, dict):
        payload = [payload]
    if not isinstance(payload, list):
        report.error("DATA_SCHEMA_SAMPLE_INVALID", f"échantillon `{sample}` : ni objet ni tableau",
                     fix="fournir un tableau d'enregistrements")
        return []
    return [r for r in payload if isinstance(r, dict)]


# ---------------------------------------------------------------------------
# Inférence
# ---------------------------------------------------------------------------
def infer_source(ctx: Context, source_id: str, report: Report, *, sample: Path | None,
                 force: bool) -> dict[str, Any] | None:
    src = ctx.registry.sources.get(source_id)
    if src is None:
        report.error("DATA_SOURCE_UNKNOWN", f"source `{source_id}` non déclarée",
                     fix="vérifier l'`id` dans `## Active Data Sources` ou dans les manifestes")
        return None

    target = ctx.schema_path(source_id)
    if target.is_file() and not force:
        report.error(
            "DATA_SCHEMA_ALREADY_FROZEN",
            f"source `{source_id}` : un schéma figé existe déjà ({paths.rel(ctx.root, target)})",
            fix="c'est voulu : le schéma figé fait foi, et le réinférer sans le dire annulerait la détection "
                "de drift. Relancer avec `--force` après avoir décidé que la donnée a légitimement changé",
        )
        return None

    records, coerce = read_records(ctx, source_id, src, report, sample=sample)
    if not records:
        if report.ok:
            report.error("DATA_SOURCE_EMPTY", f"source `{source_id}` : aucun enregistrement observé",
                         fix="un schéma inféré depuis zéro enregistrement n'est pas un schéma")
        return None

    locked = {str(src.get("key") or "")} | {str(f) for f in (src.get("required_filter") or [])}
    no_enum = {str(f) for f in (src.get("pii") or [])} | {str(f) for f in (src.get("free_text") or [])}
    result = schema_infer.infer(records, source_id=source_id, coerce=coerce,
                                string_only={f for f in locked if f}, no_enum=no_enum)

    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(result.schema, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
                      encoding="utf-8")
    for note in result.notes:
        report.warn("DATA_SCHEMA_REVIEW_REQUIRED", f"source `{source_id}` : {note}",
                    fix="corriger à la main dans le schéma figé, puis le commiter")
    report.warn(
        "DATA_SCHEMA_REVIEW_REQUIRED",
        f"source `{source_id}` : schéma écrit depuis {result.records} enregistrement(s) — À RELIRE avant de commiter",
        fix=f"relire {paths.rel(ctx.root, target)} : types, `required`, `enum`, et surtout les descriptions "
            "de champ, qui sont des gabarits et non des descriptions",
        location=paths.rel(ctx.root, target),
    )
    return result.schema


# ---------------------------------------------------------------------------
# Rendu — wrappers Python
# ---------------------------------------------------------------------------
def _py_type(spec: dict[str, Any], *, optional: bool) -> str:
    raw = spec.get("type", "string")
    names = [raw] if isinstance(raw, str) else list(raw)
    parts = [PY_TYPES.get(str(n), "Any") for n in names if str(n) != "null"]
    out = " | ".join(dict.fromkeys(parts)) or "Any"
    if optional or "null" in [str(n) for n in names]:
        out += " | None"
    return out


def _literal(value: str) -> str:
    return json.dumps(value, ensure_ascii=False)


def _field_line(name: str, spec: dict[str, Any], *, optional: bool, default: str | None = None) -> str:
    annotation = _py_type(spec, optional=optional)
    description = str(spec.get("description") or f"Champ {name} de la source.")
    args = [f"description={_literal(description)}"]
    if default is not None:
        args.insert(0, default)
    return f"    {name}: {annotation} = Field({', '.join(args)})"


def render_wrapper(ctx: Context, source_id: str, src: dict[str, Any], schema: dict[str, Any],
                   kind: str) -> str:
    props: dict[str, Any] = schema.get("properties") or {}
    required = set(schema.get("required") or [])
    key = str(src.get("key") or "")
    store_id = str(src.get("store") or "")
    connector = str(src.get("connector") or "")
    pii = sorted(str(f) for f in (src.get("pii") or []))
    untrusted = sorted(str(f) for f in (src.get("free_text") or []))
    trust = sr.source_trust(src)
    tool_name = f"{source_id}_{kind}"

    entrypoint = {"lookup": "lookup_record", "search": "search_records", "count": "count_records"}[kind]

    lines: list[str] = [
        f"SPEC = ToolSpec.from_contract({_literal(ctx.contract_id(source_id, kind))})",
        "",
        f"SOURCE = {_literal(source_id)}",
        _frozenset_line("PII_FIELDS", pii),
        _frozenset_line("UNTRUSTED_FIELDS", untrusted),
        "",
        "",
        "class Input(BaseModel):",
        '    model_config = ConfigDict(frozen=True, extra="forbid")',
        "",
    ]

    if kind == "lookup":
        spec = props.get(key) or {"type": "string", "description": f"Clé d'identification `{key}`."}
        lines.append(_field_line(key, spec, optional=False))
    else:
        lines.extend(_search_input_fields(src, props))

    lines.extend(["", "", "class Record(BaseModel):",
                  "    model_config = ConfigDict(frozen=True)", ""])
    if props:
        for name in sorted(props):
            lines.append(_field_line(name, props[name], optional=name not in required,
                                     default=None if name in required else "default=None"))
    else:
        lines.append("    pass")

    lines.extend(["", "", "class Output(BaseModel):",
                  "    model_config = ConfigDict(frozen=True)", ""])
    if kind == "lookup":
        lines.append('    record: Record | None = Field(description="L\'enregistrement trouvé, ou null si la clé est inconnue.")')
    elif kind == "search":
        lines.append('    records: list[Record] = Field(description="Enregistrements retenus, triés de façon déterministe.")')
        lines.append('    truncated: bool = Field(description="Vrai si le plafond a été atteint : le total réel est SUPÉRIEUR — affiner les filtres avant de conclure.")')
    else:
        lines.append('    count: int = Field(description="Nombre d\'enregistrements correspondant aux filtres, compté côté code et non par le modèle.")')
    lines.append('    as_of: str = Field(description="Instantané de la source, ISO 8601 UTC. La donnée a pu changer depuis.")')
    lines.append('    stale: bool = Field(description="Vrai si l\'instantané dépasse max_staleness_hours : le dire à l\'utilisateur.")')

    lines.extend(["", "", f"async def {tool_name}(params: Input, *, ctx: ToolContext) -> Output:"])
    lines.append(f'    """{_docstring(src, kind)}"""')
    if kind == "lookup":
        call = ["    return await lookup_record(", "        source=SOURCE,",
                f"        key=params.{key},", "        record_model=Record,", "        output_model=Output,"]
    elif kind == "search":
        call = ["    return await search_records(", "        source=SOURCE,",
                "        filters=params.model_dump(exclude_none=True),", "        record_model=Record,",
                "        output_model=Output,"]
    else:
        call = ["    return await count_records(", "        source=SOURCE,",
                "        filters=params.model_dump(exclude_none=True),", "        output_model=Output,"]
    call.extend(["        ctx=ctx,", "        pii_fields=PII_FIELDS,",
                 "        untrusted_fields=UNTRUSTED_FIELDS,", "    )"])
    lines.extend(call)

    header = [
        f"# {BANNER}",
        f"# Source `{source_id}` ({connector} · store `{store_id}`) · confiance : {trust}.",
        f"# Schéma figé : {paths.rel(ctx.root, ctx.schema_path(source_id))}",
        "# Éditer ce fichier est [DATA_TOOL_HAND_EDITED] : corriger la DÉCLARATION de la",
        "# source, puis `gen_source_tools.py --write`. Le code et la déclaration ne peuvent",
        "# pas diverger sans que quelqu'un s'en aperçoive — c'est tout l'intérêt.",
        "from __future__ import annotations",
        "",
    ]
    # `Any` n'apparaît que si un champ du schéma figé est un objet ou un tableau.
    # Importer un symbole inutilisé ferait échouer le lint du projet généré, et un
    # générateur dont la sortie ne passe pas le lint n'est jamais rebranché.
    if any(re.search(r"\bAny\b", line) for line in lines):
        header.extend(["from typing import Any", ""])
    header.extend([
        "from pydantic import BaseModel, ConfigDict, Field",
        "",
        f"from {ctx.app}.data.envelope import {entrypoint}",
        f"from {ctx.app}.tools.spec import ToolContext, ToolSpec",
        "",
        "#: Classe d'effet de bord de CE code — confrontée au contrat par la TOOL GATE.",
        "#: Une source déclarée est en lecture seule par construction (declared-sources).",
        'SIDE_EFFECT_CLASS = "read-only"',
        "",
    ])
    return "\n".join(header + lines) + "\n"


def _frozenset_line(name: str, values: list[str]) -> str:
    if not values:
        return f"{name}: frozenset[str] = frozenset()"
    inner = ", ".join(_literal(v) for v in values)
    return f"{name} = frozenset({{{inner}}})"


def _search_input_fields(src: dict[str, Any], props: dict[str, Any]) -> list[str]:
    required_filters = [str(f) for f in (src.get("required_filter") or [])]
    filters = [str(f) for f in (src.get("filters") or [])]
    ranges = [str(f) for f in (src.get("ranges") or [])]
    out: list[str] = []

    for name in sorted(set(required_filters)):
        spec = props.get(name) or {"type": "string"}
        out.append(_field_line(name, {**spec, "description": _filter_doc(name, spec, mandatory=True)},
                               optional=False))
    for name in sorted(set(filters) - set(required_filters)):
        spec = props.get(name) or {"type": "string"}
        out.append(_field_line(name, {**spec, "description": _filter_doc(name, spec, mandatory=False)},
                               optional=True, default="default=None"))
    for name in sorted(set(ranges)):
        spec = props.get(name) or {"type": "string"}
        for bound, word in (("min", "borne basse incluse"), ("max", "borne haute incluse")):
            out.append(_field_line(
                f"{name}_{bound}",
                {**spec, "description": f"Filtre de plage sur `{name}` : {word}. Même type et même unité que le champ."},
                optional=True, default="default=None"))
    if not out:
        out.append("    pass")
    return out


def _filter_doc(name: str, spec: dict[str, Any], *, mandatory: bool) -> str:
    head = "Filtre OBLIGATOIRE" if mandatory else "Filtre optionnel"
    enum = spec.get("enum")
    tail = f" Valeurs admises : {', '.join(str(v) for v in enum)}." if isinstance(enum, list) and enum else ""
    return f"{head} sur `{name}` — égalité stricte, jamais une expression.{tail}"


def _docstring(src: dict[str, Any], kind: str) -> str:
    first = next((l.strip() for l in str(src.get("description") or "").split("\n") if l.strip()), "")
    verbs = {"lookup": "Lecture par clé", "search": "Recherche filtrée", "count": "Comptage filtré"}
    return f"{verbs[kind]}. {first}".replace('"', "'").strip()


# ---------------------------------------------------------------------------
# Rendu — squelettes de tool-contract
# ---------------------------------------------------------------------------
def _tool_schemas(src: dict[str, Any], schema: dict[str, Any], kind: str,
                  max_rows: int) -> tuple[dict[str, Any], dict[str, Any]]:
    props: dict[str, Any] = {k: _clean(v) for k, v in (schema.get("properties") or {}).items()}
    key = str(src.get("key") or "")

    if kind == "lookup":
        param = props.get(key) or {"type": "string"}
        input_schema = {
            "type": "object", "additionalProperties": False,
            "properties": {key: {**param, "description": f"Clé d'identification de l'enregistrement recherché (`{key}`)."}},
            "required": [key],
        }
    else:
        properties: dict[str, Any] = {}
        required: list[str] = []
        for name in sorted({str(f) for f in (src.get("required_filter") or [])}):
            properties[name] = {**(props.get(name) or {"type": "string"}),
                                "description": _filter_doc(name, props.get(name) or {}, mandatory=True)}
            required.append(name)
        for name in sorted({str(f) for f in (src.get("filters") or [])} - set(required)):
            properties[name] = {**(props.get(name) or {"type": "string"}),
                                "description": _filter_doc(name, props.get(name) or {}, mandatory=False)}
        for name in sorted({str(f) for f in (src.get("ranges") or [])}):
            base = props.get(name) or {"type": "string"}
            properties[f"{name}_min"] = {**base, "description": f"Borne basse incluse sur `{name}`."}
            properties[f"{name}_max"] = {**base, "description": f"Borne haute incluse sur `{name}`."}
        input_schema = {"type": "object", "additionalProperties": False,
                        "properties": properties or {}, "required": required}

    record_schema = {"type": "object", "properties": props,
                     "required": sorted(schema.get("required") or [])}
    common = {
        "as_of": {"type": "string", "format": "date-time",
                  "description": "Instantané de la source en ISO 8601 UTC. La donnée a pu changer depuis : toute réponse temporelle doit en tenir compte."},
        "stale": {"type": "boolean",
                  "description": "Vrai si l'instantané dépasse max_staleness_hours ; l'agent doit alors le signaler à l'utilisateur."},
    }
    if kind == "lookup":
        output = {"type": "object",
                  "properties": {"record": {"type": ["object", "null"], **record_schema,
                                            "description": "L'enregistrement trouvé, ou null si la clé est inconnue."},
                                 **common},
                  "required": ["as_of", "stale"]}
    elif kind == "search":
        output = {"type": "object",
                  "properties": {"records": {"type": "array", "items": record_schema,
                                             "description": f"Au plus {max_rows} enregistrements, triés de façon déterministe."},
                                 "truncated": {"type": "boolean",
                                               "description": f"Vrai si le plafond de {max_rows} a été atteint : le total réel est SUPÉRIEUR, ne jamais conclure sur ce sous-ensemble."},
                                 **common},
                  "required": ["records", "truncated", "as_of", "stale"]}
    else:
        output = {"type": "object",
                  "properties": {"count": {"type": "integer",
                                           "description": "Nombre d'enregistrements correspondant aux filtres, compté côté code. Utiliser CET outil pour toute question « combien »."},
                                 **common},
                  "required": ["count", "as_of", "stale"]}
    return input_schema, output


def _clean(spec: Any) -> dict[str, Any]:
    if not isinstance(spec, dict):
        return {"type": "string"}
    return {k: v for k, v in spec.items() if not k.startswith("x-")}


def render_contract(ctx: Context, source_id: str, src: dict[str, Any], schema: dict[str, Any],
                    kind: str) -> str:
    max_rows = ctx.envelope_int("SourceMaxRecordsReturned", 200)
    timeout_s = max(1, math.ceil(ctx.envelope_int("SourceReadTimeoutMs", 5000) / 1000))
    store = ctx.registry.stores.get(str(src.get("store") or "")) or {}
    rpm = store.get("rate_limit_rpm") or 600
    trust = sr.source_trust(src)
    input_schema, output_schema = _tool_schemas(src, schema, kind, max_rows)
    auth_env = sorted(sr.auth_env_refs(store))
    contract_id = ctx.contract_id(source_id, kind)
    max_bytes = max_rows * 1024 if kind == "search" else 65536

    rows = "\n".join(
        f"| `{code}` | {meaning} | {behavior} |"
        for code, (meaning, behavior) in DECLARED_ERRORS[kind].items()
    )
    body = [
        f"# TOOL CONTRACT: {contract_id}",
        "",
        f"MISSION: {ctx.mission}-{ctx.mission_name.split('-', 1)[1] if '-' in ctx.mission_name else ctx.mission_name}",
        "Status: Draft",
        "Side Effect Class: read-only",
        f"Trust: {trust}",
        "",
        "> SQUELETTE généré par `gen_source_tools.py` : §1, §2, §4, §5 et §6 viennent de la",
        "> déclaration de la source. Ce fichier, lui, SE COMPLÈTE — `architect-tools` remplit",
        "> §7, `dev-prompt` revoit la description de §1, `qa-evals` déclare la suite",
        "> de §8 (`qa-tests` en écrit les tests). La régénération ne l'écrase jamais. En revanche, la description de §1 doit",
        "> rester celle de la source : un seul artefact porte l'intention métier, et",
        "> `--check` signale la divergence.",
        "",
        "---",
        "",
        "## 1. Nom et description",
        "",
        f"- **name** : `{source_id}_{kind}`",
        "- **description** :",
        "",
        "```",
        str(src.get("description") or "").rstrip(),
        "```",
        "",
        "## 2. Schémas",
        "",
        "- **Entrée** :",
        "```json",
        json.dumps(input_schema, ensure_ascii=False, indent=2, sort_keys=True),
        "```",
        "- **Sortie** :",
        "```json",
        json.dumps(output_schema, ensure_ascii=False, indent=2, sort_keys=True),
        "```",
        "",
        "## 3. Stratégie de sûreté",
        "",
        "Sans objet : `read-only`. La stack `declared-sources` est en lecture seule par",
        "construction — un besoin d'écriture est `dataaccess/repository-tools.md`.",
        "",
        "## 4. Erreurs déclarées",
        "",
        "| Code | Signification | Comportement attendu de l'agent |",
        "|---|---|---|",
        rows,
        "",
        "## 5. Bornes techniques",
        "",
        "| | |",
        "|---|---|",
        f"| `timeout_s` | {timeout_s} |",
        f"| `rate_limit_rpm` | {rpm} |",
        "| `retry_policy` | exponential:2 |",
        f"| `max_response_bytes` | {max_bytes} |",
        "",
        "## 6. Authentification",
        "",
    ]
    if auth_env:
        body.append(f"- **Variable d'environnement** : `{auth_env[0]}`")
        if len(auth_env) > 1:
            body.append(f"- Autres variables du store : {', '.join('`' + v + '`' for v in auth_env[1:])}")
        body.append("")
        body.append("La valeur vit dans le fichier de secrets gitignoré, jamais ici.")
    else:
        body.append(f"Sans objet : le store `{src.get('store')}` est déclaré `auth: {{ mode: none }}`.")
    body.extend([
        "",
        "## 7. Exposé à",
        "",
        "| Agent | CAP qui l'exige |",
        "|---|---|",
        "| `{n}-<agent>` | `{n}-{m}-<CAP>` |",
        "",
        "> À compléter par `architect-tools`. Si aucune CAP ne l'exige, l'outil ne doit",
        "> être câblé à personne (`[TOOL_SCOPE_EXCESS]`).",
        "",
        "## 8. Tests de contrat (L2)",
        "",
        f"Fichier : `workspace/proof/suites/tool-{contract_id}.yaml`",
        "",
    ])
    body.extend(f"- [ ] {item}" for item in _test_checklist(kind, src))
    if trust == "untrusted":
        body.extend([
            "",
            "## 9. Si `trust: untrusted`",
            "",
            _untrusted_note(src),
        ])
    return "\n".join(body) + "\n"


def _test_checklist(kind: str, src: dict[str, Any]) -> list[str]:
    items = ["happy path"] + [f"erreur `{code}`" for code in DECLARED_ERRORS[kind]]
    items.append("`as_of` présent dans chaque réponse, y compris vide")
    if kind == "search":
        items.append(f"plafond atteint -> `truncated: true` (lecture de maxRows+1)")
        items.append("ordre déterministe : deux appels identiques rendent la même liste")
    if kind == "count":
        items.append("le compte porte sur le TOTAL, pas sur la page tronquée")
    if src.get("pii"):
        items.append(f"champs PII redigés dans les spans : {', '.join(str(f) for f in src['pii'])}")
    if src.get("free_text"):
        items.append(f"champs de texte libre enveloppés : {', '.join(str(f) for f in src['free_text'])}")
    return items


def _untrusted_note(src: dict[str, Any]) -> str:
    free = ", ".join(f"`{f}`" for f in (src.get("free_text") or [])) or "la réponse entière"
    return (
        f"La sortie de cet outil est écrite par un tiers : {free} est traité comme du texte hostile (P8).\n"
        "Le wrapper l'enveloppe par `wrap_untrusted` avant qu'elle atteigne le contexte du modèle,\n"
        "et chaque agent qui consomme cet outil doit porter une suite d'injection indirecte —\n"
        "sans quoi une phrase déposée dans la donnée devient une instruction."
    )


# ---------------------------------------------------------------------------
# Modes
# ---------------------------------------------------------------------------
def planned(ctx: Context, report: Report, only: str | None = None) -> list[tuple[str, dict[str, Any], str]]:
    """Les (source, déclaration, kind) que la déclaration commande de générer."""
    out: list[tuple[str, dict[str, Any], str]] = []
    for source_id in sorted(ctx.registry.sources):
        if only and source_id != only:
            continue
        src = ctx.registry.sources[source_id]
        if str(src.get("connector") or "") not in sr.CONNECTORS:
            continue
        kinds = ["lookup"] if src.get("key") else []
        if src.get("filters") or src.get("ranges") or src.get("required_filter"):
            kinds.extend(["search", "count"])
        if not kinds:
            report.warn(
                "DATA_SOURCE_NO_TOOL",
                f"source `{source_id}` : ni `key` ni `filters`/`ranges` — aucun outil à générer",
                fix="déclarer au moins une clé ou un filtre : une source qu'on ne peut pas interroger "
                    "n'est pas une surface, c'est un fichier",
            )
        out.extend((source_id, src, kind) for kind in kinds)
    return out


def load_frozen(ctx: Context, source_id: str, report: Report) -> dict[str, Any] | None:
    path = ctx.schema_path(source_id)
    if not path.is_file():
        report.error(
            "DATA_SOURCE_SCHEMA_MISSING",
            f"source `{source_id}` : aucun schéma figé ({paths.rel(ctx.root, path)})",
            fix=f"`gen_source_tools.py --infer --source {source_id}`, puis RELIRE le fichier avant de commiter",
        )
        return None
    try:
        schema = json.loads(markdown_io.read_text(path))
    except ValueError as exc:
        report.error("DATA_SOURCE_SCHEMA_MALFORMED", f"source `{source_id}` : schéma figé illisible ({exc})",
                     fix="régénérer le schéma", location=paths.rel(ctx.root, path))
        return None
    if not isinstance(schema, dict) or not isinstance(schema.get("properties"), dict):
        report.error("DATA_SOURCE_SCHEMA_MALFORMED", f"source `{source_id}` : schéma figé sans `properties`",
                     fix="un schéma de source est un JSON Schema d'objet",
                     location=paths.rel(ctx.root, path))
        return None
    return schema


#: Le runtime — code INVARIANT, identique pour tout projet d'un même langage.
#:
#: Il est généré, et non écrit par un agent, pour trois raisons : c'est le même
#: code à chaque fois ; une régression s'y verrait sur tous les projets à la
#: fois ; et `--check` attrape une retouche à la main. Ce qui varie vit dans
#: `sources.json` (le registre résolu) et dans `tools/` (un wrapper par source).
RUNTIME_ROOT = "templates/runtime"

#: Les sous-arbres du runtime que CE générateur émet. Allowlist explicite, et
#: non « tout ce qui traîne sous `templates/runtime/{lang}/` » : l'arbre porte
#: aussi `app/`, le squelette applicatif, qui appartient à `gen_app_skeleton.py`
#: et se déclenche sur une autre condition. Un projet en `declared-sources` doit
#: recevoir ses outils de données, pas une application complète qu'il n'a pas
#: demandée — et que `--check` lui réclamerait ensuite indéfiniment.
RUNTIME_SUBTREES: tuple[str, ...] = ("data", "tools")


def runtime_dir(root: Path, language: str) -> Path:
    local = root / ".sdda" / RUNTIME_ROOT / language
    return local if local.is_dir() else paths.FRAMEWORK_SDDA_DIR / RUNTIME_ROOT / language


def render_registry(ctx: Context, report: Report) -> dict[str, Any]:
    """`sources.json` — le registre RÉSOLU, seul état que le runtime lit.

    L'application ne parse jamais STACK.md : c'est une déclaration de projet, pas
    une configuration d'exécution, et surtout ce qui est résolu est **épinglable** (P10). Deux runs sur
    deux registres différents ne sont pas comparables, et le pipeline doit
    pouvoir le dire plutôt qu'aligner deux scores incomparables.
    """
    section = ctx.section
    envelope = {
        "role": str(section.get("SourceAgentRole") or "readonly"),
        "read_timeout_ms": ctx.envelope_int("SourceReadTimeoutMs", 5000),
        "max_records_returned": ctx.envelope_int("SourceMaxRecordsReturned", 200),
        "max_object_bytes": ctx.envelope_int("SourceMaxObjectBytes", 52_428_800),
        "schema_check_sample": ctx.envelope_int("SourceSchemaCheckSample", 500),
        "max_staleness_hours": ctx.envelope_int("SourceMaxStalenessHours", 24),
        "forbidden_ops": _as_tuple(section.get("SourceForbiddenOps")),
        "allowed_sources": _as_tuple(section.get("SourceAllowedSources")) or tuple(sorted(ctx.registry.sources)),
        "allowed_stores": _as_tuple(section.get("SourceAllowedStores")) or tuple(sorted(ctx.registry.stores)),
        "egress_allowlist": _as_tuple(section.get("SourceEgressAllowlist")),
        "query_logging": str(section.get("SourceQueryLogging") or "full"),
    }

    stores = [{"id": sid, "kind": str(s.get("kind") or ""),
               "root": str(s["root"]) if s.get("root") else None,
               "base_url": str(s["base_url"]) if s.get("base_url") else None,
               "server": str(s["server"]) if s.get("server") else None,
               "read_only": bool(s.get("read_only", True))}
              for sid, s in sorted(ctx.registry.stores.items())]

    sources = []
    for sid, src in sorted(ctx.registry.sources.items()):
        entry: dict[str, Any] = {
            "id": sid, "connector": str(src.get("connector") or ""),
            "store": str(src.get("store") or ""), "key": str(src.get("key") or ""),
            "description": str(src.get("description") or "").strip(),
            "trust": sr.source_trust(src), "encoding": str(src.get("encoding") or "utf-8"),
            "header_row": int(src.get("header_row") or 1),
        }
        for field_name in ("glob", "format", "delimiter", "sheet", "records_path", "date_field"):
            value = src.get(field_name)
            entry[field_name] = str(value) if value not in (None, "") else None
        for field_name in ("filters", "ranges", "required_filter", "pii", "free_text"):
            entry[field_name] = _as_tuple(src.get(field_name))
        stale = src.get("max_staleness_hours")
        entry["max_staleness_hours"] = int(stale) if isinstance(stale, int) else None
        sources.append(entry)

    payload = {
        "version": 1,
        "generated_from": "workspace/stack/STACK.md + " + ", ".join(ctx.registry.manifests)
        if ctx.registry.manifests else "workspace/stack/STACK.md",
        "envelope": envelope, "stores": stores, "sources": sources,
    }
    payload["content_hash"] = "sha256:" + hashlib.sha256(
        json.dumps({k: v for k, v in payload.items() if k != "content_hash"},
                   ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    return payload


def _as_tuple(value: Any) -> tuple[str, ...]:
    if isinstance(value, (list, tuple)):
        return tuple(str(v).strip() for v in value if str(v).strip())
    return ()


def render_tool_specs(ctx: Context, report: Report, only: str | None) -> dict[str, Any]:
    """`tool_specs.json` — les contrats RÉSOLUS des outils de données.

    L'application ne lit jamais les contrats Markdown : ils ne sont pas livrés
    avec elle, et surtout `tool_schema_hash` doit désigner ce que l'outil FAIT,
    pas ce qu'un fichier disait le jour où on l'a relu. Le hash porte donc sur
    le couple (description du contrat, schémas d'entrée/sortie) — les deux
    seules choses dont un changement modifie le comportement observable.
    """
    max_rows = ctx.envelope_int("SourceMaxRecordsReturned", 200)
    timeout_s = max(1, math.ceil(ctx.envelope_int("SourceReadTimeoutMs", 5000) / 1000))
    tools: list[dict[str, Any]] = []

    for source_id, src, kind in planned(ctx, report, only):
        schema = load_frozen(ctx, source_id, report)
        if schema is None:
            continue
        input_schema, output_schema = _tool_schemas(src, schema, kind, max_rows)
        description = str(src.get("description") or "").strip()
        store = ctx.registry.stores.get(str(src.get("store") or "")) or {}
        pinned = json.dumps({"description": description, "input": input_schema,
                             "output": output_schema}, ensure_ascii=False, sort_keys=True)
        tools.append({
            "id": ctx.contract_id(source_id, kind),
            "name": f"{source_id}_{kind}",
            "description": description,
            "side_effect_class": "read-only",
            "trust": sr.source_trust(src),
            "timeout_s": timeout_s,
            "rate_limit_rpm": int(store.get("rate_limit_rpm") or 600),
            "retry_policy": "exponential:2",
            "max_response_bytes": max_rows * 1024 if kind == "search" else 65536,
            "errors": sorted(DECLARED_ERRORS[kind]),
            "tool_schema_hash": "sha256:" + hashlib.sha256(pinned.encode()).hexdigest(),
        })
    return {"version": 1, "tools": sorted(tools, key=lambda e: e["id"])}


def emit_runtime(ctx: Context, report: Report, *, write: bool,
                 only: str | None = None) -> tuple[list[str], list[str]]:
    """Copie le runtime et écrit `sources.json`. Rend `(écrits, divergents)`."""
    source_dir = runtime_dir(ctx.root, "python")
    if not source_dir.is_dir():
        report.error("DATA_RUNTIME_MISSING", f"runtime introuvable ({source_dir})",
                     fix="restaurer `.sdda/templates/runtime/python/`")
        return [], []

    written: list[str] = []
    drifted: list[str] = []
    targets: list[tuple[Path, str]] = []

    # Seuls `data/` et `tools/` : ce générateur est celui de l'ACCÈS AUX
    # SOURCES, et il ne se déclenche que sur `declared-sources`. Le squelette
    # applicatif voisin (`templates/runtime/python/app/`) est émis par
    # `gen_app_skeleton.py`, sur un autre déclencheur — un `rglob("*.py")` sur
    # tout l'arbre livrerait une application entière à un projet qui n'a demandé
    # que ses outils de données, et `--check` la réclamerait ensuite à chaque
    # passage.
    for subtree in RUNTIME_SUBTREES:
        for path in sorted((source_dir / subtree).rglob("*.py")):
            relative = path.relative_to(source_dir)
            targets.append((ctx.src_root / relative, markdown_io.read_text(path)))

    registry_json = json.dumps(render_registry(ctx, report), ensure_ascii=False,
                               indent=2, sort_keys=True) + "\n"
    targets.append((ctx.src_root / "data" / "sources.json", registry_json))
    targets.append((ctx.src_root / "data" / "tool_specs.json",
                    json.dumps(render_tool_specs(ctx, report, only), ensure_ascii=False,
                               indent=2, sort_keys=True) + "\n"))

    for target, content in targets:
        rel = paths.rel(ctx.root, target)
        if write:
            target.parent.mkdir(parents=True, exist_ok=True)
            if not target.is_file() or markdown_io.read_text(target) != content:
                target.write_text(content, encoding="utf-8")
                written.append(rel)
        elif not target.is_file() or markdown_io.read_text(target) != content:
            drifted.append(rel)
    return written, drifted


def run_generate(ctx: Context, report: Report, *, write: bool, only: str | None) -> dict[str, Any]:
    wrappers_written: list[str] = []
    contracts_written: list[str] = []
    stale: list[str] = []
    missing: list[str] = []

    for source_id, src, kind in planned(ctx, report, only):
        schema = load_frozen(ctx, source_id, report)
        if schema is None:
            continue
        wrapper = render_wrapper(ctx, source_id, src, schema, kind)
        wpath = ctx.wrapper_path(source_id, kind)
        rel = paths.rel(ctx.root, wpath)

        if write:
            wpath.parent.mkdir(parents=True, exist_ok=True)
            if not wpath.is_file() or markdown_io.read_text(wpath) != wrapper:
                wpath.write_text(wrapper, encoding="utf-8")
                wrappers_written.append(rel)
        elif not wpath.is_file():
            missing.append(rel)
        elif markdown_io.read_text(wpath) != wrapper:
            stale.append(rel)

        cpath = ctx.contract_path(source_id, kind)
        crel = paths.rel(ctx.root, cpath)
        if not cpath.is_file():
            if write:
                cpath.parent.mkdir(parents=True, exist_ok=True)
                cpath.write_text(render_contract(ctx, source_id, src, schema, kind), encoding="utf-8")
                contracts_written.append(crel)
            else:
                missing.append(crel)
        else:
            _check_contract_drift(ctx, cpath, crel, source_id, src, kind, report)

    if missing:
        report.error(
            "DATA_TOOL_MISSING",
            f"{len(missing)} artefact(s) déclaré(s) mais absent(s) : {', '.join(missing[:4])}"
            + (" …" if len(missing) > 4 else ""),
            fix="`gen_source_tools.py --write` — un outil déclaré et non généré est une CAP qui échouera "
                "au premier appel, pas au build",
        )
    runtime_written, runtime_drifted = emit_runtime(ctx, report, write=write, only=only)
    if runtime_drifted:
        report.error(
            "DATA_RUNTIME_STALE",
            f"{len(runtime_drifted)} fichier(s) de runtime absent(s) ou divergent(s) : "
            + ", ".join(runtime_drifted[:4]) + (" …" if len(runtime_drifted) > 4 else ""),
            fix="`gen_source_tools.py --write`. Le runtime est du code INVARIANT : une retouche "
                "locale se perd à la régénération suivante, et diverge en silence d'ici là",
        )

    if stale:
        report.error(
            "DATA_TOOL_HAND_EDITED",
            f"{len(stale)} wrapper(s) divergent(s) de la déclaration : {', '.join(stale[:4])}"
            + (" …" if len(stale) > 4 else ""),
            fix="ces fichiers portent « NE PAS ÉDITER » : reporter le changement dans la déclaration de la "
                "source, puis `gen_source_tools.py --write`. Sinon le code et le contrat racontent deux choses",
        )

    return {
        "mission": ctx.mission, "app": ctx.app,
        "toolsDir": paths.rel(ctx.root, ctx.tools_dir()),
        "wrappersWritten": wrappers_written, "contractsWritten": contracts_written,
        "runtimeWritten": runtime_written, "runtimeDrifted": runtime_drifted,
        "stale": stale, "missing": missing,
    }


def _check_contract_drift(ctx: Context, path: Path, rel: str, source_id: str, src: dict[str, Any],
                          kind: str, report: Report) -> None:
    """Le contrat existant dit-il encore ce que la source déclare ?

    On ne compare pas le fichier entier : le contrat est **complété** par
    `architect-tools` et `dev-prompt`, et exiger l'égalité octet pour octet
    reviendrait à interdire ce travail. On compare les trois choses qui n'ont
    pas le droit de diverger : le nom de l'outil, sa posture de confiance, et
    la description — parce qu'un seul artefact doit porter l'intention métier.
    """
    text = markdown_io.read_text(path)
    header = markdown_io.parse_header_fields(text)
    naming = markdown_io.section_body(text, "Nom et description") or ""
    name = markdown_io.strip_code(markdown_io.parse_kv_list(naming).get("name", ""))
    expected_name = f"{source_id}_{kind}"
    if name and name != expected_name:
        report.error("DATA_TOOL_CONTRACT_DRIFT", f"contrat `{rel}` : `name: {name}` au lieu de `{expected_name}`",
                     fix="l'`id` de la source nomme l'outil ; renommer l'un des deux", location=rel)

    trust = str(header.get("Trust", "")).strip().lower()
    expected_trust = sr.source_trust(src)
    if trust and trust != expected_trust and not (trust == "untrusted" and expected_trust == "trusted"):
        report.error(
            "DATA_TOOL_CONTRACT_DRIFT",
            f"contrat `{rel}` : `Trust: {trust}` alors que la source est `{expected_trust}`",
            fix="un contrat peut être PLUS méfiant que la source, jamais moins : corriger le contrat, "
                "ou déclarer les champs `free_text` de la source",
            location=rel,
        )

    fences = markdown_io.fenced_blocks(naming)
    declared = _normalize(str(src.get("description") or ""))
    written = _normalize(fences[0] if fences else "")
    if written and declared and written != declared:
        report.warn(
            "DATA_TOOL_DESCRIPTION_DRIFT",
            f"contrat `{rel}` : la description diffère de celle de la source `{source_id}`",
            fix="un seul artefact porte l'intention métier. Reporter la version revue dans la déclaration "
                "de la source (elle nourrit aussi le schéma figé et les traces), puis régénérer",
            location=rel,
        )


def _normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


# ---------------------------------------------------------------------------
# Entrée
# ---------------------------------------------------------------------------
def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="Génère les outils d'accès aux sources déclarées (0 token, aucun réseau)")
    mode = p.add_mutually_exclusive_group()
    mode.add_argument("--check", action="store_true", help="défaut : comparer sans rien écrire, exit 1 si dérive")
    mode.add_argument("--write", action="store_true", help="écrire les wrappers et créer les contrats manquants")
    mode.add_argument("--infer", action="store_true", help="inférer le schéma figé d'une source (exige --source)")
    p.add_argument("--source", default=None, help="limiter à une source, par son `id`")
    p.add_argument("--mission", default=None, help="numéro de mission ; défaut : détecté si unique")
    p.add_argument("--from-sample", dest="sample", type=Path, default=None,
                   help="fichier JSON d'échantillon, pour une source distante (http-api, mcp, store non local)")
    p.add_argument("--src-root", type=Path, default=None, help="racine du paquet applicatif généré")
    p.add_argument("--force", action="store_true", help="--infer : réécrire un schéma figé existant")
    add_common_args(p)
    return p


def run(root: Path, *, mode: str, source: str | None = None, mission: str | None = None,
        sample: Path | None = None, src_root: Path | None = None, force: bool = False) -> Report:
    report = Report(name="GEN-SOURCE-TOOLS", target=str(root))

    if not paths.stack_md_path(root).is_file():
        report.error("STACK_MISSING", "STACK.md introuvable", fix="lancer `python bootstrap.py`")
        return report

    # Écrit sans littéral de liste : `[DECLARED]` dans le source serait compté
    # comme une classe d'erreur par `sync_error_registry.py`, qui scanne les
    # formes `[CLASS]`. Un faux positif dans le registre le rend moins lu.
    strategy = active_stacks(root, "Active Data Access")
    if len(strategy) != 1 or strategy[0] != DECLARED:
        report.error(
            "DATA_ACCESS_INCONSISTENT",
            f"`## Active Data Access` n'active pas `{DECLARED}` (trouvé : {strategy or '<aucune>'})",
            fix=f"ce générateur ne sert que la stack `{DECLARED}` ; pour une base, les vues et repositories "
                "sont écrits par `dev-data` depuis les contrats",
            location="workspace/stack/STACK.md",
        )
        return report

    ctx = Context(root, mission=mission, src_root=src_root)
    for problem in ctx.registry.problems:
        (report.error if problem.severity == "error" else report.warn)(
            problem.cls, problem.message, problem.fix, problem.location or "workspace/stack/STACK.md")
    if not report.ok:
        return report

    if mode == "infer":
        if not source:
            report.error("INVALID_ARG", "`--infer` exige `--source <id>`",
                         fix="un schéma s'infère source par source, et se relit source par source")
            return report
        infer_source(ctx, source, report, sample=sample, force=force)
        return report

    if not ctx.mission:
        report.error(
            "MISSION_AMBIGUOUS",
            "numéro de mission indécidable (zéro ou plusieurs missions dans workspace/feats/missions/)",
            fix="passer `--mission {n}` — il nomme les contrats générés",
        )
        return report

    report.data.update(run_generate(ctx, report, write=(mode == "write"), only=source))
    return report


def main(argv: list[str] | None = None) -> int:
    ensure_utf8_stdout()
    args = build_parser().parse_args(argv)
    mode = "infer" if args.infer else ("write" if args.write else "check")
    report = run(resolve_root(args), mode=mode, source=args.source, mission=args.mission,
                 sample=args.sample, src_root=args.src_root, force=args.force)
    return finish(report, args)


if __name__ == "__main__":
    sys.exit(main())
